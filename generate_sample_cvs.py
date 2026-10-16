"""Tạo 10 hồ sơ giả lập, mỗi hồ sơ có Word, PDF có chữ và ảnh PNG.

Mục đích: kiểm thử đọc tệp, nhận diện kỹ năng và ưu tiên nơi muốn làm việc.
Không chứa CV thật. Các tên, quá trình học tập và dự án đều là ví dụ giả lập.
PDF được nhúng font Unicode; ảnh được render từ chính PDF, không phải ảnh minh họa.
"""
from __future__ import annotations

from dataclasses import dataclass, asdict
from html import escape
from pathlib import Path
import json
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parent
DEST = ROOT / "data/sample_cvs"


@dataclass(frozen=True)
class SampleCV:
    code: str
    slug: str
    name: str
    title: str
    domain: str
    address: str
    preferred: str
    region: str
    years: int
    skills: tuple[str, ...]
    summary: str
    activities: tuple[str, ...]
    project: str
    test_note: str
    english: bool = False


CVS = [
    SampleCV("CV01", "data_analyst_ha_noi", "Nguyễn Minh Anh", "Data Analyst", "Phân tích dữ liệu",
             "Hà Nội", "Hà Nội", "ha_noi", 1,
             ("Python", "SQL", "Pandas", "Excel", "Power BI", "Statistics"),
             "Tôi muốn ứng tuyển vị trí phân tích dữ liệu, chuyển số liệu kinh doanh thành báo cáo và đề xuất có thể thực hiện. Tôi có 1 năm thực hành làm sạch dữ liệu, xây dựng dashboard và theo dõi chỉ số bán hàng.",
             ("Làm sạch dữ liệu đơn hàng bằng Python và Pandas; kiểm tra giá trị thiếu và bản ghi trùng.",
              "Viết truy vấn SQL để tổng hợp doanh thu, phân nhóm khách hàng và đối chiếu số liệu.",
              "Thiết kế dashboard Power BI và báo cáo Excel theo tháng; giải thích biến động chỉ số."),
             "Dashboard bán lẻ giả lập: từ dữ liệu CSV đến báo cáo doanh thu, tồn kho và hiệu quả bán hàng.",
             "Ưu tiên các việc phân tích dữ liệu hoặc BI tại Hà Nội."),
    SampleCV("CV02", "backend_python_uu_tien_ha_noi", "Trần Quốc Bảo", "Backend Developer Python", "Phát triển phần mềm",
             "TP. Hồ Chí Minh", "Hà Nội", "ha_noi", 2,
             ("Python", "FastAPI", "Django", "SQL", "PostgreSQL", "REST API", "Docker", "Git"),
             "Tôi có 2 năm thực hành phát triển dịch vụ web bằng Python. Tôi muốn làm Backend Developer, thiết kế API, xây dựng cơ sở dữ liệu và triển khai ứng dụng. Hiện ở miền Nam nhưng sẵn sàng chuyển nơi làm theo nguyện vọng đã ghi.",
             ("Xây dựng REST API bằng FastAPI và Django, xử lý phân quyền và xác thực người dùng.",
              "Thiết kế PostgreSQL, tối ưu truy vấn SQL và viết kiểm thử cho chức năng đặt lịch.",
              "Đóng gói dịch vụ bằng Docker và phối hợp nhóm qua Git, kiểm tra mã trước khi tích hợp."),
             "Hệ thống đặt lịch giả lập: quản lý người dùng, lịch hẹn, phân quyền và API cho giao diện web.",
             "Ca ưu tiên nguyện vọng: địa chỉ TP. Hồ Chí Minh nhưng phải gợi ý Hà Nội trước."),
    SampleCV("CV03", "frontend_react_sai_gon", "Lê Hoàng Nam", "Frontend Developer React", "Phát triển phần mềm",
             "Sài Gòn", "Sài Gòn", "ho_chi_minh", 2,
             ("JavaScript", "TypeScript", "React", "HTML", "CSS", "REST API", "Git", "Figma"),
             "Tôi ứng tuyển Frontend Developer và có 2 năm xây dựng giao diện web tương tác. Tôi quan tâm khả năng sử dụng trên điện thoại, hiệu năng tải trang và cách phối hợp giữa thiết kế với lập trình.",
             ("Phát triển giao diện React và TypeScript cho trang tìm kiếm, danh sách và chi tiết sản phẩm.",
              "Kết nối REST API, xử lý trạng thái đang tải, lỗi kết nối và phân trang dữ liệu.",
              "Chuyển thiết kế Figma thành HTML và CSS, kiểm tra bố cục trên nhiều màn hình."),
             "Cổng đặt vé giả lập: tìm kiếm, chọn lịch, giỏ vé và giao diện thích ứng cho điện thoại.",
             "Chuẩn hóa tên Sài Gòn thành vùng Hồ Chí Minh, ưu tiên công việc Frontend/React."),
    SampleCV("CV04", "ai_engineer_da_nang", "Phạm Thu Hà", "AI Engineer", "Trí tuệ nhân tạo",
             "Đà Nẵng", "Đà Nẵng", "da_nang", 3,
             ("Python", "PyTorch", "TensorFlow", "Machine Learning", "NLP", "Transformers", "FastAPI", "Docker"),
             "Tôi có 3 năm thực hành học máy và xử lý ngôn ngữ. Tôi muốn xây dựng hệ thống AI có khả năng phục vụ người dùng, từ chuẩn bị dữ liệu đến đánh giá mô hình và triển khai dịch vụ suy luận.",
             ("Chuẩn bị dữ liệu văn bản, huấn luyện và đánh giá mô hình phân loại bằng PyTorch.",
              "Thực hiện dự án NLP với Transformers, biểu diễn ngữ nghĩa và tìm kiếm văn bản.",
              "Triển khai API suy luận bằng FastAPI, đóng gói Docker và theo dõi thời gian phản hồi."),
             "Tìm kiếm tài liệu giả lập: embedding văn bản, truy hồi nội dung và đánh giá kết quả trên bộ câu hỏi thử.",
             "Ưu tiên nhóm AI/Machine Learning tại Đà Nẵng, không trả việc ngoài vùng đã chọn."),
    SampleCV("CV05", "data_engineer_ho_chi_minh", "Võ Gia Huy", "Data Engineer", "Kỹ thuật dữ liệu",
             "TP. Hồ Chí Minh", "TP. Hồ Chí Minh", "ho_chi_minh", 3,
             ("Python", "SQL", "Spark", "Airflow", "ETL", "PostgreSQL", "AWS", "dbt"),
             "Tôi có 3 năm thực hành xây dựng nền tảng dữ liệu. Tôi muốn làm Data Engineer, phát triển luồng thu thập và biến đổi dữ liệu đáng tin cậy để phục vụ báo cáo và phân tích.",
             ("Thiết kế pipeline ETL bằng Python và Airflow, xử lý lỗi và chạy lại tác vụ an toàn.",
              "Tối ưu SQL và dùng Spark để tổng hợp dữ liệu sự kiện theo ngày.",
              "Chuẩn hóa mô hình dữ liệu với dbt, kiểm tra chất lượng và lưu trữ trên PostgreSQL/AWS."),
             "Kho dữ liệu thương mại điện tử giả lập: lịch chạy dữ liệu, bảng tổng hợp và kiểm tra dữ liệu đầu vào.",
             "Phân biệt Data Engineer với Data Analyst, dù cùng có Python và SQL."),
    SampleCV("CV06", "tester_nha_trang", "Đặng Ngọc Linh", "Tester QA Engineer", "Đảm bảo chất lượng phần mềm",
             "Nha Trang, Khánh Hòa", "Nha Trang", "khanh_hoa", 2,
             ("Python", "SQL", "Selenium", "Postman", "Jira", "Git", "Agile"),
             "Tôi ứng tuyển kiểm thử phần mềm với 2 năm thực hành kiểm thử web và API. Tôi muốn làm tại khu vực ghi trong nguyện vọng, tập trung phát hiện lỗi và xác nhận chất lượng trước khi phát hành.",
             ("Phân tích yêu cầu, viết test case và ghi nhận lỗi có bước tái hiện trên Jira.",
              "Kiểm thử REST API bằng Postman, đối chiếu dữ liệu bằng SQL.",
              "Viết kịch bản kiểm thử giao diện bằng Python và Selenium; tham gia quy trình Agile."),
             "Kiểm thử trang đặt phòng giả lập: đăng ký, tìm phòng, đặt phòng và kiểm thử hồi quy.",
             "Nhận diện Nha Trang là nhóm Nha Trang/Khánh Hòa; ưu tiên Tester/QC trong vùng."),
    SampleCV("CV07", "devops_can_tho", "Bùi Đức Thành", "DevOps Engineer", "Hạ tầng điện toán đám mây",
             "Cần Thơ", "Cần Thơ", "can_tho", 3,
             ("Linux", "Docker", "Kubernetes", "CI/CD", "AWS", "Terraform", "Python", "Git"),
             "Tôi có 3 năm thực hành vận hành hệ thống Linux và tự động hóa triển khai. Tôi mong muốn làm DevOps Engineer, giúp nhóm phát triển đưa ứng dụng lên môi trường chạy ổn định và dễ khôi phục.",
             ("Xây dựng pipeline CI/CD để chạy kiểm thử, tạo image Docker và triển khai ứng dụng.",
              "Cấu hình Kubernetes và tài nguyên AWS bằng Terraform cho môi trường thử nghiệm.",
              "Viết script Python, theo dõi log và lập quy trình xử lý sự cố hệ thống Linux."),
             "Hạ tầng web giả lập: tự động triển khai, kiểm tra sức khỏe dịch vụ và khôi phục phiên bản trước.",
             "Ưu tiên DevOps/Cloud tại Cần Thơ, đối chiếu kỹ năng hạ tầng thay vì chỉ Python."),
    SampleCV("CV08", "embedded_hai_phong", "Đỗ Hải Đăng", "Embedded Software Engineer", "Phát triển phần mềm nhúng",
             "Hải Phòng", "Hải Phòng", "hai_phong", 3,
             ("C++", "Linux", "Embedded", "Git", "Networking"),
             "Tôi ứng tuyển kỹ sư phần mềm nhúng với 3 năm thực hành lập trình C/C++ trên Linux. Tôi quan tâm thiết bị kết nối mạng, tối ưu bộ nhớ và kiểm tra độ ổn định của phần mềm thiết bị.",
             ("Phát triển module C/C++ cho thiết bị mô phỏng, xử lý giao tiếp và dữ liệu cảm biến.",
              "Gỡ lỗi chương trình Linux, theo dõi bộ nhớ và kiểm tra lỗi truy cập tài nguyên.",
              "Làm việc với Git và kiến thức Networking để kiểm tra kết nối giữa các thiết bị."),
             "Thiết bị giám sát giả lập: thu thập cảm biến, truyền dữ liệu và thử nghiệm khi mất kết nối.",
             "Ưu tiên Embedded/C++ tại Hải Phòng; kiểm tra C++ không bị đồng nhất với C hay C#."),
    SampleCV("CV09", "uiux_english_nha_trang", "Mai Anh Nguyen", "UI/UX Designer", "Product Design",
             "Da Nang", "Nha Trang", "khanh_hoa", 2,
             ("Figma", "UX", "UI", "HTML", "CSS", "Agile"),
             "I am a UI/UX Designer with 2 years of practice in user research, interface design and interactive prototypes. I want to design clear digital products and work closely with development teams.",
             ("Created user journeys, wireframes and high-fidelity interfaces in Figma.",
              "Planned usability sessions, summarized feedback and improved navigation flows.",
              "Collaborated in an Agile team and used HTML/CSS knowledge for implementation handoff."),
             "Simulated travel booking app: customer research, mobile screens, clickable prototype and usability findings.",
             "CV tiếng Anh: ưu tiên Preferred location Nha Trang thay vì Address Da Nang; nhận diện UI/UX.", True),
    SampleCV("CV10", "cybersecurity_quang_ninh_khong_co_tin", "Hoàng Tuấn Kiệt", "Cybersecurity Specialist", "An ninh mạng",
             "Quảng Ninh", "Quảng Ninh", "quang_ninh", 2,
             ("Linux", "Networking", "Security", "Python", "SQL", "Docker"),
             "Tôi có 2 năm thực hành giám sát an ninh và kiểm tra cấu hình hệ thống. Tôi ứng tuyển chuyên viên an ninh mạng, tập trung phát hiện bất thường, phân tích log và cải thiện bảo mật dịch vụ.",
             ("Phân tích log Linux và Networking để xác định dấu hiệu truy cập bất thường.",
              "Viết công cụ Python phục vụ tổng hợp sự kiện và kiểm tra cấu hình Security.",
              "Rà soát quyền truy cập SQL, kiểm tra cấu hình Docker và viết hướng dẫn khắc phục."),
             "Phòng lab an ninh giả lập: thu thập log, phân loại cảnh báo và ghi nhận hướng xử lý.",
             "Ca không có tin: catalog 500 tin hiện không có vùng Quảng Ninh. Chọn vùng này phải trả rỗng, không tự chuyển vùng."),
]


class SampleCVBuilder:
    def __init__(self, destination: Path = DEST):
        self.destination = destination
        # Font Arial có đầy đủ dấu tiếng Việt trên máy Windows hiện tại.
        font_dir = Path("C:/Windows/Fonts")
        self.font_regular = font_dir / "arial.ttf"
        self.font_bold = font_dir / "arialbd.ttf"
        for name, path in (("CVArial", self.font_regular), ("CVArialBold", self.font_bold)):
            pdfmetrics.registerFont(TTFont(name, str(path)))
        for folder in ("word", "pdf", "images"):
            (destination / folder).mkdir(parents=True, exist_ok=True)

    @staticmethod
    def blocks(cv: SampleCV):
        en = cv.english
        return [
            ("title", "Synthetic candidate resume" if en else "Hồ sơ ứng tuyển giả lập"),
            ("name", cv.name),
            ("role", cv.title),
            ("meta", ("Email: " if en else "Email: ") + cv.code.lower() + "@example.com"),
            ("body", ("Address: " if en else "Địa chỉ: ") + cv.address),
            ("body", ("Preferred location: " if en else "Nơi muốn làm việc: ") + cv.preferred),
            ("body", ("Job title: " if en else "Vị trí ứng tuyển: ") + cv.title),
            ("body", ("Professional field: " if en else "Lĩnh vực: ") + cv.domain),
            ("heading", "Professional summary" if en else "Mục tiêu nghề nghiệp"),
            ("body", cv.summary),
            ("heading", "Technical skills" if en else "Kỹ năng chuyên môn"),
            ("body", "; ".join(cv.skills)),
            ("heading", "Experience" if en else "Kinh nghiệm"),
            ("body", f"{cv.years} years of practice in simulated projects." if en else f"{cv.years} năm thực hành trong các dự án mô phỏng."),
            *[("bullet", activity) for activity in cv.activities],
            ("heading", "Selected project" if en else "Dự án tiêu biểu"),
            ("body", cv.project),
            ("heading", "Education" if en else "Học vấn"),
            ("body", "Bachelor-level IT studies at a fictional training institution. All education and work history here is synthetic."
             if en else "Chương trình cử nhân CNTT tại cơ sở đào tạo giả lập. Toàn bộ quá trình học tập và kinh nghiệm trong hồ sơ là dữ liệu thử nghiệm."),
        ]

    def word(self, cv: SampleCV, path: Path):
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(.6)
        section.left_margin = section.right_margin = Inches(.68)
        for key in ("Normal", "Title", "Heading 1", "Heading 2"):
            style = doc.styles[key]
            style.font.name = "Arial"
            style.font.color.rgb = RGBColor(0, 0, 0)
            # Template Word mặc định có đường kẻ xanh trong Title. Loại bỏ
            # khỏi OOXML thay vì che bằng màu, để bản Word đúng bố cục đã chọn.
            for border in list(style.element.iter(qn("w:pBdr"))):
                border.getparent().remove(border)
            for font_element in style.element.iter(qn("w:rFonts")):
                for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                    font_element.attrib.pop(qn("w:" + attribute), None)
        normal = doc.styles["Normal"]
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.12
        for kind, text in self.blocks(cv):
            p = doc.add_paragraph(style="Title" if kind == "title" else "Heading 2" if kind == "heading" else "Normal")
            p.paragraph_format.keep_together = True
            if kind in ("title", "name", "role", "heading"):
                p.paragraph_format.keep_with_next = True
            run = p.add_run(("- " if kind == "bullet" else "") + text)
            if kind == "title":
                run.font.size = Pt(12); run.bold = True
                p.paragraph_format.space_after = Pt(9)
            elif kind == "name":
                run.font.size = Pt(22); run.bold = True
                p.paragraph_format.space_after = Pt(3)
            elif kind == "role":
                run.font.size = Pt(12); run.bold = True
                p.paragraph_format.space_after = Pt(9)
            elif kind == "heading":
                run.font.size = Pt(11); run.bold = True
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(5)
            elif kind == "meta":
                run.font.size = Pt(10)
        doc.core_properties.author = "JobVista synthetic CV fixtures"
        doc.core_properties.title = cv.code + " " + cv.title
        footer = section.footer.paragraphs[0]
        footer.add_run(f"{cv.code} | Synthetic data for JobVista testing").font.size = Pt(8)
        doc.save(path)

    def pdf(self, cv: SampleCV, path: Path):
        styles = {
            "body": ParagraphStyle("body", fontName="CVArial", fontSize=11, leading=15.2, spaceAfter=5, textColor=colors.HexColor("#25302b")),
            "title": ParagraphStyle("title", fontName="CVArialBold", fontSize=12, leading=16, spaceAfter=10, textColor=colors.black),
            "name": ParagraphStyle("name", fontName="CVArialBold", fontSize=22, leading=27, spaceAfter=4, textColor=colors.black),
            "role": ParagraphStyle("role", fontName="CVArialBold", fontSize=12, leading=17, spaceAfter=9, textColor=colors.black),
            "meta": ParagraphStyle("meta", fontName="CVArial", fontSize=10, leading=14, spaceAfter=5),
            "heading": ParagraphStyle("heading", fontName="CVArialBold", fontSize=11, leading=15, spaceBefore=10, spaceAfter=5, keepWithNext=True, textColor=colors.black),
            "bullet": ParagraphStyle("bullet", fontName="CVArial", fontSize=11, leading=15.2, leftIndent=10, firstLineIndent=-8, spaceAfter=5),
        }
        story = [Paragraph(escape(("- " if kind == "bullet" else "") + text), styles[kind]) for kind, text in self.blocks(cv)]
        def footer(canvas, document):
            canvas.setFont("CVArial", 8)
            canvas.setFillColor(colors.HexColor("#69746c"))
            canvas.drawString(49, 24, f"{cv.code} | Synthetic data for JobVista testing")
        SimpleDocTemplate(str(path), pagesize=letter, leftMargin=49, rightMargin=49,
                          topMargin=43, bottomMargin=43, title=cv.code + " " + cv.title,
                          author="JobVista synthetic CV fixtures").build(story, onFirstPage=footer, onLaterPages=footer)

    @staticmethod
    def png(pdf_path: Path, image_path: Path):
        with pdfium.PdfDocument(str(pdf_path)) as document:
            if len(document) != 1:
                raise ValueError(f"Expected a one-page test CV: {pdf_path.name}, got {len(document)}")
            document[0].render(scale=2.5).to_pil().convert("RGB").save(image_path)

    def build(self):
        manifest = []
        for cv in CVS:
            stem = f"{cv.code}_{cv.slug}"
            paths = {"word": f"word/{stem}.docx", "pdf": f"pdf/{stem}.pdf", "image": f"images/{stem}.png"}
            self.word(cv, self.destination / paths["word"])
            self.pdf(cv, self.destination / paths["pdf"])
            self.png(self.destination / paths["pdf"], self.destination / paths["image"])
            manifest.append({**asdict(cv), "files": paths})
        (self.destination / "expected_profiles.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        rows = ["# 10 CV giả lập để thử JobVista", "", "Có 10 hồ sơ khác nhau, mỗi hồ sơ gồm 3 định dạng cùng nội dung: Word, PDF có chữ và ảnh PNG (30 tệp). Tất cả là dữ liệu giả lập; email dùng tên miền example.com, không phải liên hệ thật.", "",
                "## Cách thử", "", "Mở http://127.0.0.1:8000 → Tìm việc với CV → chọn tệp → Đọc và phân tích CV → kiểm tra nơi muốn làm → Tìm công việc phù hợp. Có thể dùng cùng mã CV ở ba định dạng để so sánh việc đọc chữ trực tiếp và OCR.", "",
                "PDF trong thư mục pdf là PDF có chữ. Ảnh PNG dùng để thử OCR. Để thử PDF scan, in/chuyển ảnh PNG sang PDF bằng trình đọc ảnh; không cần dùng CV thật.", "",
                "| Mã | Vị trí | Nơi muốn làm | Kinh nghiệm | Ca kiểm tra |", "|---|---|---|---|---|"]
        for cv in CVS:
            rows.append(f"| {cv.code} | {cv.title} | {cv.preferred} | {cv.years} năm | {cv.test_note} |")
        rows += ["", "## Khi đối chiếu kết quả", "", "- CV02 ở TP. Hồ Chí Minh nhưng muốn làm Hà Nội; vùng được gợi ý phải là Hà Nội.",
                 "- CV09 viết bằng tiếng Anh, ở Da Nang nhưng muốn làm Nha Trang; nguyện vọng phải được ưu tiên.",
                 "- CV10 là ca kết quả rỗng có chủ đích: catalog hiện không có tin Quảng Ninh. Giữ nguyên vùng Quảng Ninh để thử; hệ thống không được tự đưa việc ở nơi khác vào.",
                 "- Các CV còn lại nên có tín hiệu nghề nghiệp và kỹ năng liên quan, nhưng không bắt buộc một Job_ID luôn đứng đầu. Đây là kiểm tra chức năng, chưa phải bộ nhãn đánh giá chất lượng mô hình.",
                 "- OCR có thể sai dấu tiếng Việt hoặc bỏ sót chữ. Hãy xem nội dung đã đọc và xác nhận vùng trước khi tìm. Mức kinh nghiệm là nội dung so khớp, không phải điều kiện loại ứng viên cứng.",
                 "- File expected_profiles.json lưu kỹ năng, mã vùng và đường dẫn của mỗi hồ sơ để đối chiếu tự động.", ""]
        (self.destination / "README.md").write_text("\n".join(rows), encoding="utf-8")
        self.zip()
        print(f"Created {len(CVS)} profiles / {len(CVS)*3} CV files in {self.destination}")

    def zip(self):
        archive_path = ROOT / "outputs/10_cv_mau_jobvista.zip"
        archive_path.parent.mkdir(exist_ok=True)
        with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(self.destination.rglob("*")):
                if path.is_file():
                    archive.write(path, arcname="10_cv_mau_jobvista/" + path.relative_to(self.destination).as_posix())
        print(archive_path)


if __name__ == "__main__":
    SampleCVBuilder().build()
