"""Kiểm tra lớp API mới dùng đúng engine/dữ liệu cũ, không giả lập điểm gợi ý."""
from io import BytesIO
import importlib.util
import unittest
from urllib.parse import quote
from unittest.mock import patch

HAS_PREVIEW = all(importlib.util.find_spec(name) for name in ("fastapi", "httpx"))


@unittest.skipUnless(HAS_PREVIEW, "Install requirements-preview.txt and httpx to test the new UI API")
class WebPreviewTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        from web_preview import app, catalog
        cls.settings=patch.dict('os.environ',{'JOBVISTA_SEMANTIC':'0'})
        cls.settings.start();catalog.cache_clear()
        cls.addClassCleanup(cls.settings.stop)
        cls.addClassCleanup(catalog.cache_clear)
        cls.client = TestClient(app)

    def test_bilingual_catalog_has_same_500_records(self):
        vi = self.client.get('/api/catalog').json()
        en = self.client.get('/api/catalog?language=en').json()
        self.assertEqual(len(vi['jobs']), 500)
        self.assertEqual(vi['companies'], 387)
        self.assertEqual(en['companies'], vi['companies'])
        self.assertEqual([j['id'] for j in vi['jobs']], [j['id'] for j in en['jobs']])
        self.assertIn('Hanoi', [r['label'] for r in en['regions']])
        self.assertNotEqual(vi['domains'], en['domains'])

    def test_detail_is_real_and_similar_excludes_current(self):
        detail = self.client.get('/api/jobs/VN0004?location=ha_noi').json()
        self.assertTrue(detail['job']['description'])
        self.assertTrue(detail['job']['source'].startswith('https://huggingface.co/'))
        self.assertTrue(all(job['id'] != 'VN0004' and 'ha_noi' in job['regions'] for job in detail['similar']))
        self.assertEqual(self.client.get('/api/jobs/missing').status_code, 404)

    def test_match_equals_existing_engine_and_filters_region(self):
        from web_preview import catalog
        text = 'Backend Developer. Python, Django, FastAPI, PostgreSQL, REST API, Docker. Preferred location: Hanoi.'
        response = self.client.post('/api/cv/match', json={'text': text, 'location': 'ha_noi', 'top_n': 5})
        self.assertEqual(response.status_code, 200)
        jobs = response.json()['jobs']
        expected = catalog('vi').engine.match_cv_to_jobs(text, 5, location='ha_noi')
        self.assertEqual([job['id'] for job in jobs], list(expected['Job_ID']))
        self.assertTrue(all('ha_noi' in job['regions'] for job in jobs))
        self.assertEqual(jobs[0]['score'], round(float(expected.iloc[0]['Similarity_Score']), 4))
        self.assertEqual(len(jobs[0]['components']), 6)
        self.assertAlmostEqual(sum(jobs[0]['component_weights'].values()),1.)
        self.assertEqual(response.headers['cache-control'], 'no-store')
        self.assertIn('not_mentioned_skills',jobs[0])
        self.assertEqual(response.json()['algorithm']['mode'],'hybrid-lexical')
        empty = self.client.post('/api/cv/match', json={'text': text, 'location': 'quang_ninh'}).json()
        self.assertEqual(empty['jobs'], [])

    def test_word_upload_reads_bytes_and_identifies_location(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Backend Developer: Python, SQL, Django, Docker.')
        document.add_paragraph('Preferred location: Hanoi.')
        buffer = BytesIO()
        document.save(buffer)
        response = self.client.post('/api/cv/read', content=buffer.getvalue(), headers={
            'Content-Type': 'application/octet-stream', 'X-CV-Filename': quote('Hồ sơ.docx')})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn('Python', data['profile']['skills'])
        self.assertEqual(data['profile']['locations'], ['ha_noi'])
        self.assertFalse(data['ocr'])
        self.assertIn('Django', data['text'])
        confirmed=self.client.post('/api/cv/match',json={'text':data['text'],'experience_override':True,'experience_years':3}).json()
        self.assertEqual(confirmed['profile']['experience_years'],3)

    def test_invalid_input_and_size_limit(self):
        response = self.client.post('/api/cv/read?language=en', content=b'bad', headers={'X-CV-Filename': 'x.exe'})
        self.assertEqual(response.status_code, 400)
        self.assertIn('Choose', response.json()['detail'])
        response = self.client.post('/api/cv/read', content=b'x' * (15*1024*1024+1), headers={'X-CV-Filename': 'x.pdf'})
        self.assertEqual(response.status_code, 413)
        for payload in ({'text': '   '}, {'text': 'Python', 'location': 'anywhere'}, {'text': 'Python', 'top_n': 0}):
            self.assertEqual(self.client.post('/api/cv/match', json=payload).status_code, 422)

    def test_ui_assets_and_cv_are_not_exposed_as_static_files(self):
        self.assertEqual(self.client.get('/').status_code, 200)
        self.assertIn('text/css', self.client.get('/assets/styles.css').headers['content-type'])
        self.assertEqual(self.client.get('/assets/../data/resumes.json').status_code, 404)


if __name__ == '__main__':
    unittest.main()
