"""Kiểm thử đọc định dạng thật và lọc vùng; tệp mẫu chỉ chứa dữ liệu giả."""
from io import BytesIO
from pathlib import Path
import unittest
from zipfile import ZipFile

from PIL import Image
from docx import Document
from pypdf import PdfWriter
import pandas as pd
import numpy as np

from cv_documents import ResumeDocumentReader, CVReadError, MAX_FILE_BYTES
from resume_profile import ResumeAnalyzer
from locations import find_regions, cv_location_candidates
from generate_mock_data import MockDataGenerator
from job_recommender import JobRecommenderSystem

TEXT = "Vị trí ứng tuyển: Backend Python\nĐịa chỉ: Sài Gòn\nNơi muốn làm việc: Hà Nội\nKỹ năng: Python, FastAPI, Django, PostgreSQL, REST API, Docker\n2 năm phát triển phần mềm và dịch vụ web."


class FixedOCR:
    def read(self, image):
        return TEXT


class DocumentReaderTests(unittest.TestCase):
    def test_word_tables_and_headers_are_read(self):
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Địa chỉ: Sài Gòn"
        document.add_paragraph("Vị trí ứng tuyển: Backend Python")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Python; SQL; Django"
        table.cell(0, 1).text = "Nơi muốn làm việc: Hà Nội"
        output = BytesIO()
        document.save(output)
        result = ResumeDocumentReader().read(output.getvalue(), "cv.docx")
        profile = ResumeAnalyzer().analyze(result.text)
        self.assertFalse(result.used_ocr)
        self.assertIn("SQL", profile.skills)
        self.assertEqual(profile.location_candidates, ("ha_noi",))

    def test_image_and_scanned_pdf_use_ocr(self):
        image = Image.new("RGB", (500, 700), "white")
        png, pdf = BytesIO(), BytesIO()
        image.save(png, format="PNG")
        image.save(pdf, format="PDF")
        for filename, content in [("cv.png", png.getvalue()), ("cv.pdf", pdf.getvalue())]:
            result = ResumeDocumentReader(ocr=FixedOCR()).read(content, filename)
            self.assertTrue(result.used_ocr)
            self.assertEqual(ResumeAnalyzer().analyze(result.text).location_candidates, ("ha_noi",))

    def test_invalid_and_large_files_fail_without_logging_cv(self):
        for content, filename, code in [(b"", "cv.pdf", "cv_file_empty"),
                                         (b"broken document", "cv.docx", "cv_file_invalid"),
                                         (b"broken image", "cv.png", "cv_file_invalid"),
                                         (b"a" * (MAX_FILE_BYTES + 1), "cv.pdf", "cv_file_large"),
                                         (b"old Word", "cv.doc", "cv_format_unsupported")]:
            with self.subTest(filename=filename, code=code), self.assertRaises(CVReadError) as caught:
                ResumeDocumentReader().read(content, filename)
            self.assertEqual(caught.exception.code, code)

    def test_pdf_page_limit_and_password_are_reported(self):
        writer = PdfWriter()
        for _ in range(21):
            writer.add_blank_page(width=400, height=600)
        content = BytesIO()
        writer.write(content)
        with self.assertRaises(CVReadError) as caught:
            ResumeDocumentReader().read(content.getvalue(), "cv.pdf")
        self.assertEqual(caught.exception.code, "cv_too_many_pages")
        writer = PdfWriter()
        writer.add_blank_page(width=400, height=600)
        writer.encrypt("test-password")
        content = BytesIO()
        writer.write(content)
        with self.assertRaises(CVReadError) as caught:
            ResumeDocumentReader().read(content.getvalue(), "cv.pdf")
        self.assertEqual(caught.exception.code, "cv_pdf_locked")


class ProfileAndRegionTests(unittest.TestCase):
    def test_preferred_location_overrides_home_and_old_employer(self):
        text = TEXT + "\nCông việc trước đây ở Đà Nẵng"
        profile = ResumeAnalyzer().analyze(text)
        self.assertEqual(profile.location_candidates, ("ha_noi",))
        self.assertEqual(profile.role_codes, ("backend",))
        self.assertNotIn("Địa chỉ", profile.matching_text)
        self.assertNotIn("Nơi muốn làm việc", profile.matching_text)
        self.assertEqual(find_regions("TP.HCM, Nha Trang, Hà Nội"), ["ho_chi_minh", "khanh_hoa", "ha_noi"])

    def test_unknown_preference_does_not_fall_back_to_home(self):
        self.assertEqual(cv_location_candidates("Địa chỉ: Sài Gòn\nNơi muốn làm việc: chưa rõ"), [])
        self.assertEqual(cv_location_candidates("Địa chỉ: Sài Gòn\nNơi muốn làm việc: Hà Ni", tolerant=True), ["ha_noi"])
        self.assertEqual(cv_location_candidates("Địa chỉ: Sài Gòn\nNơi muốn làm việc: Hà Ni"), [])

    def test_skills_do_not_confuse_java_and_javascript(self):
        profile = ResumeAnalyzer().analyze("Frontend Developer\nJavaScript, C#, C++, .NET, Power BI")
        self.assertNotIn("Java", profile.skills)
        self.assertIn("JavaScript", profile.skills)
        self.assertIn("C#", profile.skills)
        self.assertIn("C++", profile.skills)

    def test_region_is_hard_filter_and_weighted_scores_are_exact(self):
        jobs = MockDataGenerator().create_jobs()
        jobs["Region_Codes"] = "ho_chi_minh"
        jobs.loc[jobs["Job_ID"].isin(["J009", "J010"]), "Region_Codes"] = "ha_noi"
        engine = JobRecommenderSystem(jobs, backend="tfidf")
        results = engine.match_cv_to_jobs(TEXT, 10, location="ha_noi")
        self.assertEqual(results.iloc[0]["Job_ID"], "J009")
        self.assertTrue(results["Region_Codes"].eq("ha_noi").all())
        expected = 0.25 * results.Score_Title + 0.35 * results.Score_Skills + 0.15 * results.Score_Domain + 0.25 * results.Score_Job_Description
        np.testing.assert_allclose(results.Similarity_Score, expected)
        self.assertIn("Python", results.iloc[0]["Matched_Skills"])
        self.assertTrue(engine.match_cv_to_jobs(TEXT, location="khanh_hoa").empty)
        similar = engine.recommend_similar_jobs("J009", 10, location="ha_noi")
        self.assertEqual(similar.Job_ID.tolist(), ["J010"])
        with self.assertRaises(ValueError):
            engine.match_cv_to_jobs(TEXT, location="fake_region")

    def test_metadata_survives_validation_and_query_does_not_mutate_catalog(self):
        jobs = MockDataGenerator().create_jobs()
        jobs["Region_Codes"] = "ha_noi"
        jobs["Company"] = "Example Company"
        engine = JobRecommenderSystem(jobs, backend="tfidf")
        before = engine.jobs
        result = engine.match_cv_to_jobs(TEXT, location="ha_noi")
        self.assertIn("Company", result)
        pd.testing.assert_frame_equal(before, engine.jobs)


if __name__ == "__main__":
    unittest.main()
